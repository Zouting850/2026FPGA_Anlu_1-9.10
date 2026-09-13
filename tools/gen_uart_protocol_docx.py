#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Render README section 7 (串口屏控制) into a standalone Word handoff document.

The docx is a *derived* artifact: every table and sentence is parsed out of
README.md at run time, so it cannot drift from the doc of record. Re-run this
after editing README section 7.

    python tools/gen_uart_protocol_docx.py [-o doc/串口屏控制协议_淘晶驰.docx]
"""

import argparse
import datetime
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

CJK = "微软雅黑"
MONO = "Consolas"

HEADING = re.compile(r"^(#{3,4})\s+(.*)$")
BULLET = re.compile(r"^(\s*)-\s+(.*)$")
FENCE = re.compile(r"^\s*```")
TSEP = re.compile(r"^\s*\|[\s:|-]+\|\s*$")
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]*`)")

PIPE_ESC = "\x00"


def set_cjk(obj):
    """Point w:eastAsia at a CJK font; python-docx only sets w:ascii by default."""
    rpr = obj.element.get_or_add_rPr() if hasattr(obj, "element") else obj._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), CJK)


def _emit(par, text, size, bold, code):
    for li, piece in enumerate(text.split("\n")):
        run = par.add_run(piece)
        run.bold = bold
        run.font.name = MONO if code else CJK
        run._element.get_or_add_rPr().find(qn("w:rFonts")).set(qn("w:eastAsia"), CJK)
        if size:
            run.font.size = size
        if li < len(text.split("\n")) - 1:
            run.add_break()


def add_runs(par, text, size=None, bold=False, code=False):
    """Walk markdown inline markup into Word runs; bold and code may nest."""
    for chunk in INLINE.split(text.replace("\\|", PIPE_ESC)):
        if not chunk:
            continue
        if len(chunk) > 4 and chunk.startswith("**") and chunk.endswith("**"):
            add_runs(par, chunk[2:-2], size, True, code)
        elif len(chunk) > 2 and chunk.startswith("`") and chunk.endswith("`"):
            add_runs(par, chunk[1:-1], size, bold, True)
        else:
            _emit(par, chunk.replace(PIPE_ESC, "|").replace("\\", ""), size, bold, code)
    return par


def plain_len(cell):
    """Approximate rendered width of the widest line; CJK glyphs count double."""
    t = re.sub(r"\*\*|`", "", cell.replace(PIPE_ESC, "|"))
    return max(sum(2 if ord(ch) > 0x2E80 else 1 for ch in ln) for ln in t.split("\n"))


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def repeat_header(row):
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def split_row(line):
    body = line.strip()
    body = body[1:] if body.startswith("|") else body
    body = body[:-1] if body.endswith("|") else body
    cells = [c.strip().replace(" ⏎ ", "\n") for c in body.replace("\\|", PIPE_ESC).split("|")]
    return cells


def add_table(doc, rows, usable_cm, body_pt):
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    widths = [max(plain_len(c) for c in [r[i] for r in rows]) for i in range(ncols)]
    widths = [max(w, 6) for w in widths]
    total = sum(widths)
    cm = [max(usable_cm * w / total, 1.1) for w in widths]
    scale = usable_cm / sum(cm)
    cm = [c * scale for c in cm]

    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblpr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)

    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci, text in enumerate(row):
            cell = cells[ci]
            cell.width = Cm(cm[ci])
            par = cell.paragraphs[0]
            par.paragraph_format.space_before = Pt(1)
            par.paragraph_format.space_after = Pt(1)
            add_runs(par, text, size=Pt(body_pt), bold=(ri == 0))
            if ri == 0:
                shade(cell, "DCE6F1")
            elif ri % 2 == 0:
                shade(cell, "F4F6FA")
    repeat_header(table.rows[0])
    doc.add_paragraph()
    return table


def extract_section(md_lines, want):
    """Return the lines of the '### <want>.' section, heading excluded.

    '####' subheadings belong to the section; only the next '##' or '###' ends it.
    """
    out, active = [], False
    for line in md_lines:
        if re.match(r"^#{2,3}\s", line):
            if active:
                break
            if re.match(r"^###\s+%s\." % re.escape(want), line):
                active = True
                continue
        if active:
            out.append(line)
    return out


def render(doc, lines, usable_cm):
    i, body_pt = 0, 8.5
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue

        if FENCE.match(line):
            i += 1
            while i < len(lines) and not FENCE.match(lines[i]):
                par = doc.add_paragraph()
                par.paragraph_format.space_before = Pt(0)
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.left_indent = Cm(0.6)
                run = par.add_run(lines[i].rstrip())
                run.font.name = MONO
                run.font.size = Pt(9)
                i += 1
            i += 1
            doc.add_paragraph()
            continue

        if line.lstrip().startswith("|") and i + 1 < len(lines) and TSEP.match(lines[i + 1]):
            rows = [split_row(line)]
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows, usable_cm, body_pt)
            continue

        m = HEADING.match(line)
        if m:
            doc.add_heading(re.sub(r"\*\*|`", "", m.group(2)), level=max(len(m.group(1)) - 2, 1))
            i += 1
            continue

        m = BULLET.match(line)
        if m:
            depth = min(len(m.group(1)) // 2, 2)
            par = doc.add_paragraph(style=["List Bullet", "List Bullet 2", "List Bullet 3"][depth])
            par.paragraph_format.space_before = Pt(1)
            par.paragraph_format.space_after = Pt(1)
            add_runs(par, m.group(2), size=Pt(10))
            i += 1
            continue

        if line.lstrip().startswith(">"):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.6)
            par.paragraph_format.space_before = Pt(3)
            par.paragraph_format.space_after = Pt(3)
            for run in add_runs(par, line.lstrip()[1:].strip(), size=Pt(9.5)).runs:
                run.italic = True
            i += 1
            continue

        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(3)
        par.paragraph_format.space_after = Pt(3)
        add_runs(par, line.strip(), size=Pt(10))
        i += 1


def main():
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    ap = argparse.ArgumentParser()
    ap.add_argument("--readme", default=os.path.join(root, "README.md"))
    ap.add_argument("-o", "--out", default=os.path.join(root, "doc", "串口屏控制协议_淘晶驰.docx"))
    ap.add_argument("--section", default="7")
    args = ap.parse_args()

    with open(args.readme, encoding="utf-8") as fh:
        md = fh.read().splitlines()

    lines = extract_section(md, args.section)
    if not lines:
        sys.exit("README section %s not found in %s" % (args.section, args.readme))

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for attr in ("left_margin", "right_margin"):
        setattr(sec, attr, Cm(1.4))
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    usable = 29.7 - 2.8

    normal = doc.styles["Normal"]
    normal.font.name = CJK
    normal.font.size = Pt(10)
    set_cjk(normal)
    for lvl, pt in (("Heading 1", 18), ("Heading 2", 15), ("Heading 3", 12.5)):
        st = doc.styles[lvl]
        st.font.name = CJK
        st.font.size = Pt(pt)
        set_cjk(st)

    title = doc.add_heading("串口屏控制协议（淘晶驰 TJC / Nextion）", level=0)
    for run in title.runs:
        run.font.name = CJK
        set_cjk(run)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("2026 安路赛题一 · 基于 EG4S20 的 HDMI 多媒体播放系统 · 生成于 %s"
                    % datetime.date.today().isoformat())
    r.font.size = Pt(10)
    r.font.name = CJK
    set_cjk(r)

    note = doc.add_paragraph()
    r = note.add_run(
        "本文件由 tools/gen_uart_protocol_docx.py 从 README.md 第 %s 节自动渲染，"
        "表格与文字全部逐字取自 README，不含独立维护的副本。"
        "协议若有改动，改 README 后重跑该脚本即可，不要直接编辑本 docx。"
        "README 表格里表示「两行分开写」的 ⏎ 标记，在本文中已换成单元格内的真实换行。" % args.section)
    r.font.size = Pt(9)
    r.italic = True
    r.font.name = CJK
    set_cjk(r)

    render(doc, lines, usable)
    doc.save(args.out)
    tables = len(doc.tables)
    print("wrote %s  (%d tables, %.1f KB)" % (args.out, tables, os.path.getsize(args.out) / 1024))


if __name__ == "__main__":
    main()
